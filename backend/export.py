"""Veri analizi için dışa aktarma (CSV / Excel / JSON).

Veri setleri:
  - history  : her fiyat kaydı bir satır (uzun/"tidy" format — pandas, Excel pivot)
  - products : ürün başına özet (ilk/son görülme, kayıt sayısı, min/maks/ort/son fiyat)
  - deals    : şu anki doğrulanmış fırsatlar
  - all      : üçü birden (xlsx'te 3 sayfa, json'da tek nesne, csv'de zip)

Sayılar makine-okunur tutulur (ondalık nokta, ISO tarih). CSV UTF-8 BOM'lu
yazılır ki Excel Türkçe karakterleri doğru göstersin; Excel'de doğrudan
çalışmak için xlsx önerilir (TR Excel CSV'de ';' ayırıcı bekler).
"""
from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
import zipfile
from datetime import datetime
from typing import Optional

from src.core.weight import parse_weight
from src.scrapers.base import BaseScraper
from src.storage.db import PriceDatabase

from . import services
from .schemas import deal_to_dict

DATASETS = ("history", "products", "deals", "all")
FORMATS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}


def _connect() -> sqlite3.Connection:
    # Salt-okunur bağlantı: tarama sırasında dışa aktarma yazımı engellemez.
    # DB hiç yoksa (ilk kurulum / CI) şemayı oluştur ki boş dosya dönsün.
    if not os.path.exists(services.db_path()):
        PriceDatabase(services.db_path()).close()
    uri = "file:" + services.db_path().replace("\\", "/") + "?mode=ro"
    return sqlite3.connect(uri, uri=True)


def _since_clause(days: Optional[int]) -> tuple[str, list]:
    if days is None or days <= 0:
        return "", []
    return " AND ph.recorded_at >= datetime('now', ?)", [f"-{int(days)} days"]


def history_rows(days: Optional[int] = None) -> list[dict]:
    where, params = _since_clause(days)
    sql = f"""
        SELECT ph.recorded_at, p.site, p.brand, p.name, p.url,
               ph.price, ph.original_price, ph.currency
        FROM price_history ph JOIN products p ON p.url = ph.product_url
        WHERE 1=1 {where}
        ORDER BY ph.recorded_at, p.site, p.brand
    """
    with _connect() as c:
        rows = c.execute(sql, params).fetchall()
    out = []
    for rec, site, brand, name, url, price, orig, cur in rows:
        page_disc = round((1 - price / orig) * 100, 2) if orig and orig > price else None
        w = parse_weight(name)
        out.append({
            "recorded_at": str(rec)[:19],
            "scan_date": str(rec)[:10],
            "site": site,
            "brand": brand,
            "product_type": BaseScraper._classify_product(name) or "",
            "name": name,
            "price": price,
            "original_price": orig,
            "page_discount_pct": page_disc,
            "weight_g": w["total_g"] if w else None,
            "price_per_kg": round(price / w["total_g"] * 1000, 2) if w else None,
            "currency": cur,
            "url": url,
        })
    return out


def product_rows(days: Optional[int] = None) -> list[dict]:
    where, params = _since_clause(days)
    sql = f"""
        SELECT p.url, p.site, p.brand, p.name, p.first_seen,
               MAX(ph.recorded_at), COUNT(*), MIN(ph.price), MAX(ph.price), AVG(ph.price),
               (SELECT price FROM price_history x WHERE x.product_url = p.url
                ORDER BY x.recorded_at DESC LIMIT 1)
        FROM products p JOIN price_history ph ON ph.product_url = p.url
        WHERE 1=1 {where}
        GROUP BY p.url
        ORDER BY p.site, p.brand, p.name
    """
    with _connect() as c:
        rows = c.execute(sql, params).fetchall()
    out = []
    for url, site, brand, name, first, last, n, mn, mx, avg, latest in rows:
        w = parse_weight(name)
        out.append({
            "site": site,
            "brand": brand,
            "product_type": BaseScraper._classify_product(name) or "",
            "name": name,
            "first_seen": str(first)[:19],
            "last_seen": str(last)[:19],
            "weight_g": w["total_g"] if w else None,
            "weight_label": w["label"] if w else None,
            "records": n,
            "latest_price": latest,
            "min_price": mn,
            "max_price": mx,
            "avg_price": round(avg, 2) if avg is not None else None,
            "latest_vs_min_pct": round((latest / mn - 1) * 100, 2) if mn else None,
            "url": url,
        })
    return out


def deal_rows() -> list[dict]:
    _, deals = services.snapshot(services.load_config())
    keys = ["verdict", "site", "brand", "name", "weight_g", "price_per_kg", "price",
            "original_price", "claimed_pct",
            "ref_low_30d", "discount_pct", "badge_seen", "spikes", "historical_min",
            "previous_logged_price", "history_points", "reason", "currency", "url"]
    return [{k: deal_to_dict(d).get(k) for k in keys} for d in deals]


def collect(dataset: str, days: Optional[int]) -> dict[str, list[dict]]:
    if dataset == "history":
        return {"history": history_rows(days)}
    if dataset == "products":
        return {"products": product_rows(days)}
    if dataset == "deals":
        return {"deals": deal_rows()}
    return {"history": history_rows(days), "products": product_rows(days), "deals": deal_rows()}


# --- yazıcılar ---------------------------------------------------------------
def _csv_bytes(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def _xlsx_bytes(sheets: dict[str, list[dict]]) -> bytes:
    from openpyxl import Workbook  # lazy: yalnızca xlsx istenince

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        if not rows:
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for r in rows:
            ws.append([r[h] for h in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(headers, start=1):
            width = max(len(h), *(len(str(r[h] or "")) for r in rows[:200]))
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width + 2, 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build(dataset: str, fmt: str, days: Optional[int]) -> tuple[bytes, str, str]:
    """(içerik, media_type, dosya adı) döndürür."""
    sheets = collect(dataset, days)
    stamp = datetime.now().strftime("%Y%m%d")
    base = f"coffee_{dataset}_{stamp}"

    if fmt == "xlsx":
        return _xlsx_bytes(sheets), FORMATS["xlsx"], base + ".xlsx"
    if fmt == "json":
        payload = sheets if dataset == "all" else sheets[dataset]
        data = json.dumps(payload, ensure_ascii=False, indent=1).encode("utf-8")
        return data, FORMATS["json"], base + ".json"
    if dataset == "all":
        zbuf = io.BytesIO()
        with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, rows in sheets.items():
                z.writestr(f"coffee_{name}_{stamp}.csv", _csv_bytes(rows))
        return zbuf.getvalue(), "application/zip", base + ".zip"
    return _csv_bytes(sheets[dataset]), FORMATS["csv"], base + ".csv"
